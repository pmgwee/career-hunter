from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
PAYLOAD_PATH = ROOT / "data" / "payloads" / "035-shell-graduate-program-2027-cv.json"
OUTPUT_PATH = ROOT / "output" / "docx" / "Gwee_Per_Ming_Shell_Graduate_Program_2027_R205046_Autofill.docx"


def first(mapping: dict, *keys: str, default=""):
    for key in keys:
        value = mapping.get(key)
        if value not in (None, "", []):
            return value
    return default


def contact_string(value) -> str:
    if isinstance(value, dict):
        return clean(first(value, "url", "display", default=""))
    return clean(value)


def clean(value) -> str:
    """Keep the résumé text ASCII-friendly for conservative ATS parsing."""
    if value is None:
        return ""
    text = str(value)
    return (
        text.replace("—", "-")
        .replace("–", "-")
        .replace("‑", "-")
        .replace("≤", "<=")
        .replace("≥", ">=")
        .replace("’", "'")
        .replace("“", '"')
        .replace("”", '"')
        .replace("•", "-")
    )


def set_cell_shading(*_args, **_kwargs):
    # Kept out deliberately: this document uses no tables or decorative blocks.
    return None


def set_keep_with_next(paragraph, enabled=True):
    ppr = paragraph._p.get_or_add_pPr()
    element = ppr.find(qn("w:keepNext"))
    if enabled and element is None:
        ppr.append(OxmlElement("w:keepNext"))
    elif not enabled and element is not None:
        ppr.remove(element)


def configure_styles(document: Document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.2)
    normal.paragraph_format.space_after = Pt(2.2)
    normal.paragraph_format.line_spacing = 1.0

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = None
    title.paragraph_format.space_after = Pt(1)
    title.paragraph_format.line_spacing = 1.0
    # Word's built-in Title style can carry a colored bottom rule; remove it
    # so the upload version stays strictly plain and ATS-safe.
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h1.font.size = Pt(11.5)
    h1.font.bold = True
    h1.font.color.rgb = None
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(3)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.line_spacing = 1.0

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h2.font.size = Pt(10.6)
    h2.font.bold = True
    h2.font.color.rgb = None
    h2.paragraph_format.space_before = Pt(4)
    h2.paragraph_format.space_after = Pt(1)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.line_spacing = 1.0

    # A compact bullet style keeps descriptions readable without introducing tables.
    if "ATS Bullet" not in styles:
        bullet = styles.add_style("ATS Bullet", WD_STYLE_TYPE.PARAGRAPH)
    else:
        bullet = styles["ATS Bullet"]
    bullet.base_style = normal
    bullet.font.name = "Calibri"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    bullet.font.size = Pt(10.0)
    bullet.paragraph_format.left_indent = Inches(0.22)
    bullet.paragraph_format.first_line_indent = Inches(-0.14)
    bullet.paragraph_format.space_after = Pt(1.5)
    bullet.paragraph_format.line_spacing = 1.0


def add_paragraph(document, text="", *, style="Normal", bold=False, italic=False, size=None,
                  space_before=None, space_after=None, keep_next=False):
    paragraph = document.add_paragraph(style=style)
    if text:
        run = paragraph.add_run(clean(text))
        run.bold = bold
        run.italic = italic
        if size is not None:
            run.font.size = Pt(size)
    if space_before is not None:
        paragraph.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        paragraph.paragraph_format.space_after = Pt(space_after)
    if keep_next:
        set_keep_with_next(paragraph)
    return paragraph


def add_bullet(document, text):
    p = document.add_paragraph(style="ATS Bullet")
    p.add_run("- ").bold = False
    p.add_run(clean(text))
    p.paragraph_format.space_after = Pt(2.2)
    return p


def add_section_heading(document, text):
    return add_paragraph(document, clean(text).upper(), style="Heading 1", keep_next=True)


def add_contact_line(document, values):
    p = document.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(1)
    p.add_run(" | ".join(clean(v) for v in values if v))
    return p


def add_education(document, item):
    degree = first(item, "degree", "title", "program", "name")
    institution = first(item, "institution", "school", "company", "org")
    location = first(item, "location", "city")
    dates = first(item, "dates", "date", "period", "year")
    if " | " in clean(dates) and not location:
        dates, location = clean(dates).split(" | ", 1)
    detail = first(item, "detail", "details", "result")
    add_paragraph(document, degree, style="Heading 2", keep_next=True)
    add_paragraph(document, institution, keep_next=True)
    if location:
        add_paragraph(document, f"Location: {location}", keep_next=True)
    if dates:
        add_paragraph(document, f"Dates: {dates}", keep_next=True)
    if detail:
        add_paragraph(document, detail, space_after=1)


def add_experience(document, item):
    role = first(item, "role", "title", "job_title", "position")
    company = first(item, "company", "employer", "organization", "name")
    location = first(item, "location", "city")
    dates = first(item, "dates", "date", "period")
    bullets = first(item, "bullets", "achievements", "description", default=[])
    add_paragraph(document, role, style="Heading 2", keep_next=True)
    add_paragraph(document, company, keep_next=True)
    add_paragraph(document, f"Location: {location}", keep_next=True)
    add_paragraph(document, f"Dates: {dates}", keep_next=True)
    for bullet in bullets:
        add_bullet(document, bullet)


def add_project(document, item):
    title = first(item, "title", "name", "project")
    technologies = first(item, "technologies", "tech", "stack", default=[])
    url = first(item, "url", "link", "website")
    bullets = first(item, "bullets", "achievements", "description", default=[])
    add_paragraph(document, title, style="Heading 2", keep_next=True, space_before=6, space_after=2)
    if url:
        add_paragraph(document, f"URL: {url}", keep_next=True)
    if technologies:
        if isinstance(technologies, list):
            technologies = ", ".join(clean(value) for value in technologies)
        add_paragraph(document, f"Technologies: {technologies}", keep_next=True)
    for bullet in bullets:
        add_bullet(document, bullet)


def main():
    payload = json.loads(PAYLOAD_PATH.read_text(encoding="utf-8"))
    candidate = payload.get("candidate", payload.get("profile", {}))
    document = Document()
    configure_styles(document)

    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    name = first(candidate, "name", "full_name")
    p = add_paragraph(document, name, style="Title", keep_next=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    contact_values = [
        first(candidate, "phone", "mobile"),
        first(candidate, "email"),
        first(candidate, "location", "city"),
    ]
    add_contact_line(document, contact_values)
    links = [
        contact_string(first(candidate, "linkedin", "linkedin_url")),
        contact_string(first(candidate, "github", "github_url")),
        contact_string(first(candidate, "portfolio", "website", "portfolio_url")),
    ]
    add_contact_line(document, links)

    summary = first(payload, "summary", "professional_summary")
    if summary:
        add_section_heading(document, "Professional Summary")
        add_paragraph(document, summary)

    education = payload.get("education", [])
    if education:
        add_section_heading(document, "Education")
        for item in education:
            add_education(document, item)

    skills = payload.get("skills", {})
    if skills:
        add_section_heading(document, "Skills")
        if isinstance(skills, dict):
            for category, values in skills.items():
                if isinstance(values, list):
                    values = ", ".join(clean(value) for value in values)
                add_paragraph(document, f"{category}: {values}")
        else:
            for value in skills:
                if isinstance(value, dict):
                    category = first(value, "category", "name")
                    items = first(value, "items", "values", default=[])
                    if isinstance(items, list):
                        items = ", ".join(clean(item) for item in items)
                    add_paragraph(document, f"{category}: {items}")
                else:
                    add_paragraph(document, value)

    experience = payload.get("experience", payload.get("work_experience", []))
    if experience:
        add_section_heading(document, "Work Experience")
        for item in experience:
            add_experience(document, item)

    projects = payload.get("projects", [])
    if projects:
        add_section_heading(document, "Projects")
        for item in projects:
            add_project(document, item)

    awards = payload.get("awards", payload.get("awards_and_honors", []))
    if awards:
        add_section_heading(document, "Awards and Honors")
        for item in awards:
            if isinstance(item, dict):
                text = first(item, "title", "name", "award")
                issuer = first(item, "issuer", "organization", "by", "org")
                year = first(item, "year", "date")
                suffix = " - ".join(part for part in (issuer, year) if part)
                add_bullet(document, f"{text} - {suffix}" if suffix else text)
            else:
                add_bullet(document, item)

    volunteering = payload.get("volunteering", payload.get("volunteer", []))
    if volunteering:
        add_section_heading(document, "Volunteering")
        for item in volunteering:
            if isinstance(item, dict):
                title = first(item, "title", "role", "name")
                organization = first(item, "organization", "company", "employer")
                dates = first(item, "dates", "date", "period")
                bullets = first(item, "bullets", "description", default=[])
                add_paragraph(document, title, style="Heading 2", keep_next=True)
                if organization:
                    add_paragraph(document, organization, keep_next=True)
                if dates:
                    add_paragraph(document, f"Dates: {dates}", keep_next=True)
                for bullet in bullets:
                    add_bullet(document, bullet)
            else:
                add_bullet(document, item)

    document.core_properties.author = clean(name)
    document.core_properties.title = f"{clean(name)} Resume"
    document.core_properties.subject = "Shell Graduate Program 2027 application resume"
    document.core_properties.keywords = "resume, CV, Shell, autofill"
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
